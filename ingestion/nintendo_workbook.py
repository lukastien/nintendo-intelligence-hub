"""Parsers for Nintendo's historical Excel disclosures.

Fiscal-year columns are the flow for that April-March year. The life-to-date
column is a different span. Open-ended headers and the August fiscal calendar
are rejected. Blank cells stay missing.
"""

from __future__ import annotations

import hashlib
import re
from collections.abc import Iterable
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from urllib.parse import urljoin

from ingestion.classify import classify_period_header, month_end, normalize_label
from ingestion.numbers import parse_number
from ingestion.records import ParsedDocument, Rejection, StagedRow, dedupe_rows
from ingestion.vocabulary import INCOME_METRICS, component_platform, family_platform, region
from nih.errors import GrainError, NihError
from nih.periods import (
    DisclosureMeasure,
    PeriodSpan,
    ProductGrain,
    expected_period_bounds,
)
from nih.provenance import SourceClass

_FILE_AS_OF = re.compile(r"e(\d{2})(\d{2})\.(?:xlsx|xls)$", re.IGNORECASE)
_HREF = re.compile(
    r'href="([^"]+\.xls[x]?)"[^>]*>\s*As of\s+([A-Za-z]+)\s+(\d{1,2}),\s+(\d{4})',
    re.IGNORECASE | re.DOTALL,
)
_MONTHS = {
    "january": 1,
    "february": 2,
    "march": 3,
    "april": 4,
    "may": 5,
    "june": 6,
    "july": 7,
    "august": 8,
    "september": 9,
    "october": 10,
    "november": 11,
    "december": 12,
}
_UNIT_SCALE = 10_000
_YEN_SCALE = Decimal(1_000_000)


class WorkbookParseError(NihError):
    """Raised when a workbook cannot be mapped without guessing."""


def discover_historical_files(html: str, page_url: str) -> list[dict[str, str]]:
    """Return file URL, kind, and as-of date printed beside each historical link."""

    from datetime import date

    found: list[dict[str, str]] = []
    for href, month_name, day, year in _HREF.findall(html):
        month = _MONTHS.get(month_name.casefold())
        if month is None:
            raise WorkbookParseError(f"Unrecognized as-of month on the historical index: {month_name}")
        as_of = date(int(year), month, int(day))
        url = urljoin(page_url, href)
        name = url.rsplit("/", 1)[-1]
        suffix = _FILE_AS_OF.search(name)
        if suffix is None:
            raise WorkbookParseError(f"Historical file name has no eYYMM suffix: {name}")
        file_as_of = month_end(2000 + int(suffix.group(1)), int(suffix.group(2)))
        if file_as_of != as_of:
            raise WorkbookParseError(
                f"{name} encodes {file_as_of.isoformat()} but the link says {as_of.isoformat()}."
            )
        lowered = name.casefold()
        if "sales" in lowered:
            kind = "sales"
        elif "pl" in lowered or "income" in lowered:
            kind = "income"
        else:
            raise WorkbookParseError(f"Historical file is neither sales nor income: {name}")
        found.append({"url": url, "kind": kind, "as_of": as_of.isoformat(), "filename": name})
    if not found:
        raise WorkbookParseError("The historical index has no dated Excel links.")
    return found


def as_of_from_filename(filename: str):
    from datetime import date

    match = _FILE_AS_OF.search(filename)
    if match is None:
        raise WorkbookParseError(f"{filename} has no eYYMM suffix, so the as-of date is unknown.")
    return date(2000 + int(match.group(1)), int(match.group(2)), month_end(2000 + int(match.group(1)), int(match.group(2))).day)


def parse_sales_workbook(
    path: Path,
    *,
    source_url: str,
    as_of,
    retrieved_at: datetime,
    raw_filename: str,
) -> ParsedDocument:
    import openpyxl

    raw = path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    document = _document(
        source_id="nintendo_historical_region_sales",
        document_id=f"nintendo_sales_transition_{as_of.strftime('%Y%m%d')}",
        title="Consolidated sales transition by region",
        source_url=source_url,
        as_of=as_of,
        period_label="Fiscal-year columns and life-to-date, units in ten thousands",
        raw_filename=raw_filename,
        digest=digest,
        retrieved_at=retrieved_at,
    )
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        numeric = 0
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            numeric += _parse_sales_sheet(document, sheet.title, rows, as_of)
    finally:
        workbook.close()
    if numeric == 0:
        raise WorkbookParseError("The sales workbook has period headers but no numeric cells.")
    dedupe_rows(document)
    return document


def parse_income_workbook(
    path: Path,
    *,
    source_url: str,
    as_of,
    retrieved_at: datetime,
    raw_filename: str,
) -> ParsedDocument:
    import xlrd

    raw = path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    book = xlrd.open_workbook(path.as_posix())
    grid = [
        [book.sheet_by_index(0).cell_value(row, column) for column in range(book.sheet_by_index(0).ncols)]
        for row in range(book.sheet_by_index(0).nrows)
    ]
    return parse_income_grid(
        grid,
        source_url=source_url,
        as_of=as_of,
        retrieved_at=retrieved_at,
        raw_filename=raw_filename,
        content_sha256=digest,
    )


def parse_income_grid(
    grid: list[list[object]],
    *,
    source_url: str,
    as_of,
    retrieved_at: datetime,
    raw_filename: str,
    content_sha256: str,
) -> ParsedDocument:
    document = _document(
        source_id="nintendo_historical_income",
        document_id=f"nintendo_income_transition_{as_of.strftime('%Y%m%d')}",
        title="Consolidated statements of income transition",
        source_url=source_url,
        as_of=as_of,
        period_label="Fiscal-year columns, millions of yen",
        raw_filename=raw_filename,
        digest=content_sha256,
        retrieved_at=retrieved_at,
    )
    header_index, headers = _period_header(grid)
    notes = _note_text(grid)
    if "fy3/1990" in notes and "7 month" in notes:
        for column, printed in list(headers.items()):
            if printed.fiscal_year_ending == 1990 and printed.period_span is PeriodSpan.FISCAL_YEAR:
                headers[column] = _rejected_copy(printed, "seven_month_stub")
    numeric = 0
    for row in grid[header_index + 1 :]:
        label = _line_label(row)
        if not label:
            continue
        folded = normalize_label(label)
        if "ratio" in folded or folded.startswith("("):
            numeric += _count_numbers(row, headers)
            document.rejections.append(
                Rejection(
                    document.source_document_id,
                    "not_a_level",
                    "Ratio rows are not yen amounts and are not loaded.",
                    period_label=label,
                )
            )
            continue
        metric = INCOME_METRICS.get(folded)
        if metric is None:
            skipped = _count_numbers(row, headers)
            numeric += skipped
            if skipped:
                document.rejections.append(
                    Rejection(
                        document.source_document_id,
                        "unmapped_metric",
                        "The printed line has no metric code. It was not given a substitute code.",
                        period_label=label,
                        printed_value=str(skipped),
                    )
                )
            continue
        for column, printed in headers.items():
            if column >= len(row):
                continue
            number = parse_number(row[column])
            if number is None:
                if printed.accepted:
                    document.missing_cells += 1
                continue
            numeric += 1
            if not printed.accepted:
                document.rejections.append(
                    Rejection(
                        document.source_document_id,
                        printed.reason or "rejected_period",
                        "The column is not a March 31 fiscal year.",
                        period_label=printed.verbatim,
                        printed_value=format(number, "f"),
                    )
                )
                continue
            start, end = expected_period_bounds(printed.fiscal_year_ending, PeriodSpan.FISCAL_YEAR)  # type: ignore[arg-type]
            try:
                measure = DisclosureMeasure(
                    source_document_id=document.source_document_id,
                    source_class=SourceClass.OFFICIAL_IR,
                    metric_code=metric,
                    product_grain=ProductGrain.FINANCIAL_LINE,
                    platform_code="CONSOLIDATED",
                    region_code="GLOBAL",
                    fiscal_year_ending=printed.fiscal_year_ending,
                    period_span=PeriodSpan.FISCAL_YEAR,
                    period_start=start,
                    period_end=end,
                    as_of_date=as_of,
                    value=number * _YEN_SCALE,
                    is_additive_flow=True,
                is_region_aggregate=True,
                sales_basis="not_applicable",
            )
            except GrainError as exc:
                document.rejections.append(
                    Rejection(
                        document.source_document_id,
                        "grain",
                        str(exc),
                        period_label=printed.verbatim,
                        printed_value=format(number, "f"),
                    )
                )
                continue
            document.rows.append(
                StagedRow(
                    measure=measure,
                    printed_value=number,
                    scale_token="millions",
                    period_label_verbatim=printed.verbatim,
                    platform_name_verbatim="Consolidated",
                    region_name_verbatim="Worldwide",
                    game_title=None,
                    sheet_name="income",
                    source_url=source_url,
                    content_sha256=content_sha256,
                    retrieved_at=retrieved_at,
                    definition_note=(
                        f"Column header {printed.verbatim}. Printed scale is millions of yen. "
                        "Consolidated Nintendo Co., Ltd. statement. Not a quarterly column."
                    ),
                )
            )
    if numeric == 0:
        raise WorkbookParseError("The income sheet has no numeric cells.")
    dedupe_rows(document)
    return document


def _parse_sales_sheet(document: ParsedDocument, sheet_name: str, grid: list[list[object]], as_of) -> int:
    if not grid:
        return 0
    try:
        header_index, headers = _period_header(grid)
    except WorkbookParseError:
        document.rejections.append(
            Rejection(
                document.source_document_id,
                "unrecognized_sheet",
                "No fiscal-year or life-to-date header was found.",
                period_label=sheet_name,
            )
        )
        return 0
    scale_note = _note_text(grid)
    if "ten thousand" not in scale_note:
        raise WorkbookParseError(f"{sheet_name} does not say the figures are in ten thousands.")
    first_period = min(headers)
    europe_is_split = _sheet_has_region(grid, header_index, first_period, "europe")
    family_code = ""
    family_name = ""
    family_aggregate = False
    family_blocked = False
    component_code = ""
    component_name = ""
    component_parent = ""
    component_blocked = False
    metric = ""
    region_code = ""
    region_name = ""
    region_aggregate = False
    region_note = ""
    numeric = 0
    for row in grid[header_index + 1 :]:
        label_cells = list(enumerate(row[:first_period]))
        texts = [(index, cell) for index, cell in label_cells if isinstance(cell, str) and cell.strip()]
        row_family = ""
        row_component = ""
        row_metric = ""
        row_region = ""
        for index, cell in texts:
            folded = normalize_label(cell)
            if folded in {"hardware", "software"}:
                row_metric = folded
                continue
            if folded.startswith("of which "):
                row_component = cell
                continue
            if region(cell, europe_is_split=europe_is_split) is not None:
                row_region = cell
                continue
            if folded.startswith("released on") or folded.startswith("the figures") or folded.startswith("the breakdown"):
                continue
            if folded.startswith("[note]") or folded.startswith("(note)"):
                continue
            if index == 0:
                row_family = cell
        if row_family:
            mapped = family_platform(row_family)
            family_name = " ".join(row_family.split())
            component_code = ""
            component_name = ""
            component_parent = ""
            component_blocked = False
            if mapped is None:
                family_code = ""
                family_aggregate = False
                family_blocked = True
            else:
                family_code, family_aggregate = mapped
                family_blocked = False
        if row_component:
            component_name = " ".join(row_component.split())
            mapped_component = component_platform(row_component)
            if mapped_component is None:
                component_code = ""
                component_parent = ""
                component_blocked = True
            else:
                component_code, component_parent = mapped_component
                component_blocked = False
        elif row_metric:
            component_code = ""
            component_name = ""
            component_parent = ""
            component_blocked = False
        if row_metric:
            metric = "hw_units_sell_in" if row_metric == "hardware" else "sw_units_sell_in"
        if row_region:
            mapped_region = region(row_region, europe_is_split=europe_is_split)
            if mapped_region is not None:
                region_code, region_aggregate, region_note = mapped_region
                region_name = " ".join(row_region.split())
        if family_blocked or component_blocked:
            _reject_row_numbers(
                document,
                row,
                headers,
                reason="unmapped_platform",
                detail="The printed platform name has no warehouse code.",
                platform_name=component_name or family_name,
                region_name=region_name,
            )
            numeric += len(_numbers_in_period_columns(row, headers))
            continue
        using_component = bool(component_code) and metric == "hw_units_sell_in"
        if using_component and component_parent != family_code:
            _reject_row_numbers(
                document,
                row,
                headers,
                reason="component_parent_mismatch",
                detail=f"{component_name} is not a component of {family_name or 'the current platform'}.",
                platform_name=component_name,
                region_name=region_name,
            )
            continue
        platform_code = component_code if using_component else family_code
        platform_name = component_name if using_component else family_name
        if not platform_code or not metric or not region_code:
            orphan_numbers = _numbers_in_period_columns(row, headers)
            if orphan_numbers and row_family and not family_code:
                for printed_period, number in orphan_numbers:
                    document.rejections.append(
                        Rejection(
                            document.source_document_id,
                            "unmapped_platform",
                            "The printed platform name has no warehouse code.",
                            period_label=printed_period.verbatim,
                            platform_name=platform_name or family_name,
                            region_name=region_name,
                            printed_value=format(number, "f"),
                        )
                    )
                    numeric += 1
            continue
        for column, printed_period in headers.items():
            if column >= len(row):
                continue
            number = parse_number(row[column])
            if number is None:
                if printed_period.accepted:
                    document.missing_cells += 1
                continue
            numeric += 1
            if not printed_period.accepted:
                document.rejections.append(
                    Rejection(
                        document.source_document_id,
                        printed_period.reason or "rejected_period",
                        "The column is not a single fiscal year or a life-to-date total.",
                        period_label=printed_period.verbatim,
                        platform_name=platform_name,
                        region_name=region_name,
                        printed_value=format(number, "f"),
                    )
                )
                continue
            try:
                document.rows.append(
                    _sales_row(
                        document,
                        sheet_name=sheet_name,
                        as_of=as_of,
                        printed_period=printed_period,
                        number=number,
                        metric=metric,
                        platform_code=platform_code,
                        platform_name=platform_name,
                        aggregate=family_aggregate and not using_component,
                        component=using_component,
                        parent=component_parent if using_component else None,
                        region_code=region_code,
                        region_name=region_name,
                        region_aggregate=region_aggregate,
                        region_note=region_note,
                    )
                )
            except GrainError as exc:
                document.rejections.append(
                    Rejection(
                        document.source_document_id,
                        "grain",
                        str(exc),
                        period_label=printed_period.verbatim,
                        platform_name=platform_name,
                        region_name=region_name,
                        printed_value=format(number, "f"),
                    )
                )
    return numeric


def _sales_row(
    document: ParsedDocument,
    *,
    sheet_name: str,
    as_of,
    printed_period,
    number: Decimal,
    metric: str,
    platform_code: str,
    platform_name: str,
    aggregate: bool,
    component: bool,
    parent: str | None,
    region_code: str,
    region_name: str,
    region_aggregate: bool,
    region_note: str,
) -> StagedRow:
    if printed_period.period_span is PeriodSpan.LIFE_TO_DATE:
        fiscal_year = None
        start = None
        end = as_of
        quarter = None
    else:
        fiscal_year = printed_period.fiscal_year_ending
        quarter = printed_period.fiscal_quarter
        start, end = expected_period_bounds(fiscal_year, printed_period.period_span, quarter)
    grain = ProductGrain.HARDWARE if metric == "hw_units_sell_in" else ProductGrain.SOFTWARE_TOTAL
    note = (
        f"Column header {printed_period.verbatim}. Sheet {sheet_name}. "
        "Printed unit is ten thousand. The fiscal-year column is that year's sell-in, "
        "not a quarterly total and not life-to-date."
    )
    if printed_period.period_span is PeriodSpan.LIFE_TO_DATE:
        note = (
            f"Column header {printed_period.verbatim}. Sheet {sheet_name}. "
            "Printed unit is ten thousand. Life-to-date is not a fiscal-year flow."
        )
    if region_note:
        note = f"{note} {region_note}"
    measure = DisclosureMeasure(
        source_document_id=document.source_document_id,
        source_class=SourceClass.OFFICIAL_IR,
        metric_code=metric,
        product_grain=grain,
        platform_code=platform_code,
        region_code=region_code,
        fiscal_year_ending=fiscal_year,
        period_span=printed_period.period_span,
        fiscal_quarter=quarter,
        period_start=start,
        period_end=end,
        as_of_date=as_of,
        value=number * _UNIT_SCALE,
        is_platform_aggregate=aggregate,
        is_platform_component=component,
        parent_platform_code=parent,
        is_region_aggregate=region_aggregate,
        sales_basis="sell_in",
    )
    return StagedRow(
        measure=measure,
        printed_value=number,
        scale_token=str(_UNIT_SCALE),
        period_label_verbatim=printed_period.verbatim,
        platform_name_verbatim=platform_name,
        region_name_verbatim=region_name,
        game_title=None,
        sheet_name=sheet_name,
        source_url=document.source_url,
        content_sha256=document.content_sha256,
        retrieved_at=document.retrieved_at,
        definition_note=note,
    )


def _document(
    *,
    source_id: str,
    document_id: str,
    title: str,
    source_url: str,
    as_of,
    period_label: str,
    raw_filename: str,
    digest: str,
    retrieved_at: datetime,
) -> ParsedDocument:
    return ParsedDocument(
        source_id=source_id,
        source_document_id=document_id,
        title=title,
        source_url=source_url,
        publisher="Nintendo Co., Ltd.",
        published_on=as_of,
        period_label_verbatim=period_label,
        raw_filename=raw_filename,
        content_sha256=digest,
        retrieved_at=retrieved_at,
    )


def _period_header(grid: Iterable[list[object]]):
    for index, row in enumerate(grid):
        headers = {}
        for column, cell in enumerate(row):
            if not isinstance(cell, str):
                continue
            printed = classify_period_header(cell)
            if printed is not None:
                headers[column] = printed
        if headers:
            return index, headers
    raise WorkbookParseError("No period header row was found.")


def _note_text(grid: Iterable[list[object]]) -> str:
    parts: list[str] = []
    for row in grid:
        for cell in row:
            if isinstance(cell, str) and cell.strip():
                parts.append(normalize_label(cell))
    return "\n".join(parts)


def _line_label(row: list[object]) -> str:
    for cell in row[:3]:
        if isinstance(cell, str) and cell.strip():
            return " ".join(cell.split())
    return ""


def _count_numbers(row: list[object], headers: dict) -> int:
    count = 0
    for column in headers:
        if column < len(row) and parse_number(row[column]) is not None:
            count += 1
    return count


def _sheet_has_region(grid, header_index: int, first_period: int, folded_region: str) -> bool:
    for row in grid[header_index + 1 :]:
        for cell in row[:first_period]:
            if isinstance(cell, str) and normalize_label(cell) == folded_region:
                return True
    return False


def _numbers_in_period_columns(row, headers):
    found = []
    for column, printed in headers.items():
        if column < len(row):
            number = parse_number(row[column])
            if number is not None:
                found.append((printed, number))
    return found


def _reject_row_numbers(document, row, headers, *, reason, detail, platform_name, region_name) -> None:
    for printed, number in _numbers_in_period_columns(row, headers):
        document.rejections.append(
            Rejection(
                document.source_document_id,
                reason,
                detail,
                period_label=printed.verbatim,
                platform_name=platform_name,
                region_name=region_name,
                printed_value=format(number, "f"),
            )
        )


def _rejected_copy(printed, reason: str):
    from ingestion.classify import PrintedPeriod

    return PrintedPeriod(printed.verbatim, False, reason=reason)
